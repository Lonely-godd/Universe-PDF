import shutil
import subprocess
import zipfile
from unittest.mock import MagicMock

import pytest
from pypdf import PdfReader

from universe_pdf.conversion import ConversionError, Converter


def test_missing_office(tmp_path, font):
    path = tmp_path / "file.docx"
    path.write_bytes(b"test")
    with pytest.raises(ConversionError, match="LibreOffice"):
        Converter(font_path=font, office_binary="missing-soffice-binary").compile(
            [path], tmp_path / "out.pdf"
        )


@pytest.mark.office
@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice не установлен")
def test_real_docx(tmp_path, font):
    path = tmp_path / "document.docx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr(
            "[Content_Types].xml",
            """<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>""",
        )
        archive.writestr(
            "_rels/.rels",
            """<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>""",
        )
        archive.writestr(
            "word/document.xml",
            """<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Тест документа</w:t></w:r></w:p></w:body></w:document>""",
        )
    output = tmp_path / "result.pdf"
    Converter(font_path=font).compile([path], output)
    assert len(PdfReader(output).pages) >= 1


@pytest.mark.office
@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice не установлен")
@pytest.mark.parametrize("extension", ["xlsx", "pptx"])
def test_real_spreadsheet_and_slides(tmp_path, font, extension):
    path = tmp_path / f"document.{extension}"
    if extension == "xlsx":
        from openpyxl import Workbook

        book = Workbook()
        book.active.append(["Документ", "Количество"])
        book.active.append(["Фотографии", 3])
        book.save(path)
    else:
        from pptx import Presentation

        book = Presentation()
        slide = book.slides.add_slide(book.slide_layouts[0])
        slide.shapes.title.text = "Общий документ"
        book.save(path)
    output = tmp_path / "result.pdf"
    Converter(font_path=font).compile([path], output)
    reader = PdfReader(output)
    assert len(reader.pages) >= 1
    assert any("документ" in page.extract_text().lower() for page in reader.pages)


@pytest.mark.parametrize("timeout", [True, False])
def test_office_timeout_or_missing_output(tmp_path, font, monkeypatch, timeout):
    path = tmp_path / "file.docx"
    path.write_bytes(b"test")
    monkeypatch.setattr("universe_pdf.conversion.shutil.which", lambda value: "soffice")
    process = MagicMock()
    process.__enter__.return_value = process
    process.returncode = 0
    if timeout:
        process.wait.side_effect = [subprocess.TimeoutExpired("soffice", 1), None]
    monkeypatch.setattr("universe_pdf.conversion.subprocess.Popen", lambda *a, **k: process)
    monkeypatch.setattr("universe_pdf.conversion.subprocess.run", lambda *a, **k: None)
    with pytest.raises(ConversionError, match="не успел" if timeout else "не смог"):
        Converter(font_path=font).compile([path], tmp_path / "out.pdf")
    if timeout:
        process.kill.assert_called_once()
